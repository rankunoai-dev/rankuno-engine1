"""Tests for the deliverables HTTP surface (cycle 0087).

Covers the two things the Step 5 audit pre-step flagged as non-negotiable:
org-scoping (every read/status/download enforces `record.org_id == org_id`,
mirroring `get_job`/`get_result`) and async execution (nothing here blocks a
request handler on `build_workbook`; every build-triggering endpoint returns
`202` and the work finishes on a worker thread this test polls for).
"""

from __future__ import annotations

import io
import time
import zipfile

import openpyxl
import pytest
from fastapi.testclient import TestClient
from src.api.server import API_PREFIX, create_app
from src.core.config import get_settings
from src.core.schemas import OrgConfig
from src.core.state_store import DiskJobStore, DiskOrgConfigStore
from src.core.url_safety import UrlSafetyPolicy
from src.modules.seo.deliverables.rulebook import RULEBOOK_SHEET_NAME
from src.modules.seo.page_classifier.discovery import DiscoveryReport
from src.modules.seo.page_classifier.schemas import (
    ConsensusMethod,
    DiscoverySource,
    FullPageIntelligenceProfile,
    HierarchyLevel,
    Indexability,
    PrimaryPageType,
    SearchIntent,
    SignalScore,
    SignalSource,
)
from src.modules.seo.page_classifier.tool import CrawlSummary, PageClassificationOutput
from src.modules.seo.page_classifier.weights import SiteProfile, WeightProfileReport

from tests.api.conftest import TEST_SESSION_SECRET, auth_headers

PUBLIC_IP = "93.184.216.34"
SAFE_URL = "https://e.com/"


def profile(url: str) -> FullPageIntelligenceProfile:
    """A minimal valid profile — just enough for `to_audit_dataset` to accept it."""
    return FullPageIntelligenceProfile(
        url=url,
        canonical_url=url,
        normalized_path=url,
        hierarchy_level=HierarchyLevel.L3_LEAF_PAGE,
        primary_page_type=PrimaryPageType.UNKNOWN,
        depth_from_l0=1,
        search_intent=SearchIntent.INFORMATIONAL,
        signals_evaluated=(
            SignalScore(
                source=SignalSource.SITEMAP_INDEX,
                suggested_level=HierarchyLevel.L3_LEAF_PAGE,
                suggested_page_type=PrimaryPageType.UNKNOWN,
                confidence=0.5,
            ),
        ),
        final_confidence_score=0.5,
        consensus_method=ConsensusMethod.LAYER1_STRUCTURAL,
        discovery_sources=DiscoverySource(sitemap=True, dom_link=True),
        indexability=Indexability.INDEXABLE,
        indexability_reason="",
    )


def fake_crawl_output() -> PageClassificationOutput:
    return PageClassificationOutput(
        base_url=SAFE_URL,
        site_profile=SiteProfile(),
        weight_profile=WeightProfileReport.for_site(SiteProfile()),
        discovery=DiscoveryReport(base_url=SAFE_URL),
        summary=CrawlSummary(pages_classified=1),
        pages=(profile(SAFE_URL),),
    )


def write_rulebook_bytes() -> bytes:
    buffer = io.BytesIO()
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = RULEBOOK_SHEET_NAME
    sheet.append(("URL Pattern", "Rule Type", "Theme 1"))
    sheet.append(("/", "starts with", "Home"))
    book.save(buffer)
    return buffer.getvalue()


def write_sf_bundle_bytes() -> bytes:
    """A minimal Screaming Frog export zip: just the spine file, one row."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "internal_all.csv",
            f"Address,Content Type,Status Code,Indexability\n{SAFE_URL},text/html,200,Indexable\n",
        )
    return buffer.getvalue()


@pytest.fixture
def crawl_store(tmp_path) -> DiskJobStore:
    return DiskJobStore(tmp_path / "jobs")


@pytest.fixture
def org_store(tmp_path) -> DiskOrgConfigStore:
    store = DiskOrgConfigStore(tmp_path / "orgs")
    store.create(OrgConfig(org_id="org-a", display_name="Org A", is_active=True))
    store.create(OrgConfig(org_id="org-b", display_name="Org B", is_active=True))
    return store


@pytest.fixture(autouse=True)
def mock_org_store(monkeypatch, org_store):
    settings = get_settings()
    monkeypatch.setattr(settings, "_org_config_store", org_store)


@pytest.fixture
def client(tmp_path, crawl_store):
    app = create_app(
        store=crawl_store,
        url_policy=UrlSafetyPolicy(resolver=lambda host: [PUBLIC_IP]),
        deliverable_jobs_root=tmp_path / "deliverable_jobs",
        rulebooks_root=tmp_path / "rulebooks",
        session_secret=TEST_SESSION_SECRET,
    )
    with TestClient(app, headers=auth_headers()) as test_client:
        yield test_client


def finished_crawl_job(crawl_store: DiskJobStore, org_id: str = "org-a") -> str:
    record = crawl_store.create("seo.page_classifier", {"base_url": SAFE_URL}, org_id=org_id)
    crawl_store.finish(record.id, fake_crawl_output().model_dump(mode="json"))
    return record.id


def poll_deliverable(client, deliverable_id: str, org_id: str = "org-a") -> dict:
    deadline = time.monotonic() + 5.0
    while time.monotonic() < deadline:
        body = client.get(
            f"{API_PREFIX}/deliverables/{deliverable_id}", headers={"X-Org-Id": org_id}
        ).json()
        if body["status"] in ("succeeded", "partial", "failed"):
            return body
        time.sleep(0.01)
    pytest.fail(f"deliverable {deliverable_id} never reached a terminal status")


class TestRulebookUpload:
    def test_a_valid_upload_is_accepted_and_listed(self, client):
        response = client.post(
            f"{API_PREFIX}/deliverables/rulebooks?label=English",
            content=write_rulebook_bytes(),
            headers={"X-Org-Id": "org-a"},
        )
        assert response.status_code == 201, response.text
        body = response.json()
        assert body["org_id"] == "org-a"
        assert body["label"] == "English"
        assert body["rule_count"] == 1

        listed = client.get(
            f"{API_PREFIX}/deliverables/rulebooks", headers={"X-Org-Id": "org-a"}
        ).json()
        assert [r["id"] for r in listed] == [body["id"]]

    def test_an_empty_body_is_400(self, client):
        response = client.post(
            f"{API_PREFIX}/deliverables/rulebooks", content=b"", headers={"X-Org-Id": "org-a"}
        )
        assert response.status_code == 400

    def test_unreadable_content_is_400(self, client):
        response = client.post(
            f"{API_PREFIX}/deliverables/rulebooks",
            content=b"not a workbook",
            headers={"X-Org-Id": "org-a"},
        )
        assert response.status_code == 400

    def test_listing_is_org_scoped(self, client):
        client.post(
            f"{API_PREFIX}/deliverables/rulebooks",
            content=write_rulebook_bytes(),
            headers={"X-Org-Id": "org-a"},
        )
        listed_b = client.get(
            f"{API_PREFIX}/deliverables/rulebooks", headers={"X-Org-Id": "org-b"}
        ).json()
        assert listed_b == []

    def test_deleting_another_orgs_rulebook_is_403(self, client):
        created = client.post(
            f"{API_PREFIX}/deliverables/rulebooks",
            content=write_rulebook_bytes(),
            headers={"X-Org-Id": "org-a"},
        ).json()

        response = client.delete(
            f"{API_PREFIX}/deliverables/rulebooks/{created['id']}", headers={"X-Org-Id": "org-b"}
        )
        assert response.status_code == 403

    def test_deleting_an_unknown_rulebook_is_404(self, client):
        response = client.delete(
            f"{API_PREFIX}/deliverables/rulebooks/nope", headers={"X-Org-Id": "org-a"}
        )
        assert response.status_code == 404

    def test_a_deleted_rulebook_is_gone(self, client):
        created = client.post(
            f"{API_PREFIX}/deliverables/rulebooks",
            content=write_rulebook_bytes(),
            headers={"X-Org-Id": "org-a"},
        ).json()

        response = client.delete(
            f"{API_PREFIX}/deliverables/rulebooks/{created['id']}", headers={"X-Org-Id": "org-a"}
        )
        assert response.status_code == 204
        listed = client.get(
            f"{API_PREFIX}/deliverables/rulebooks", headers={"X-Org-Id": "org-a"}
        ).json()
        assert listed == []


class TestBuildFromJob:
    def test_an_unknown_source_job_is_404(self, client):
        response = client.post(f"{API_PREFIX}/jobs/nope/deliverable", headers={"X-Org-Id": "org-a"})
        assert response.status_code == 404

    def test_another_orgs_source_job_is_403(self, client, crawl_store):
        job_id = finished_crawl_job(crawl_store, org_id="org-a")
        response = client.post(
            f"{API_PREFIX}/jobs/{job_id}/deliverable", headers={"X-Org-Id": "org-b"}
        )
        assert response.status_code == 403

    def test_an_unfinished_source_job_is_409(self, client, crawl_store):
        record = crawl_store.create("seo.page_classifier", {"base_url": SAFE_URL}, org_id="org-a")
        response = client.post(
            f"{API_PREFIX}/jobs/{record.id}/deliverable", headers={"X-Org-Id": "org-a"}
        )
        assert response.status_code == 409

    def test_a_build_from_a_finished_job_succeeds_and_downloads(self, client, crawl_store):
        job_id = finished_crawl_job(crawl_store)

        accepted = client.post(
            f"{API_PREFIX}/jobs/{job_id}/deliverable", headers={"X-Org-Id": "org-a"}
        )
        assert accepted.status_code == 202, accepted.text
        deliverable_id = accepted.json()["id"]

        status_body = poll_deliverable(client, deliverable_id)
        assert status_body["status"] == "succeeded"
        assert status_body["has_result"] is True

        download = client.get(
            f"{API_PREFIX}/deliverables/{deliverable_id}/download", headers={"X-Org-Id": "org-a"}
        )
        assert download.status_code == 200
        assert download.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "attachment" in download.headers["content-disposition"]
        # A real workbook, not an empty blob — openpyxl can load it back.
        book = openpyxl.load_workbook(io.BytesIO(download.content))
        assert "Pages" in book.sheetnames

    def test_an_unknown_rulebook_id_is_404_before_any_build_starts(self, client, crawl_store):
        job_id = finished_crawl_job(crawl_store)
        response = client.post(
            f"{API_PREFIX}/jobs/{job_id}/deliverable",
            json={"rulebook_id": "nope"},
            headers={"X-Org-Id": "org-a"},
        )
        assert response.status_code == 404

    def test_another_orgs_rulebook_id_is_403(self, client, crawl_store):
        job_id = finished_crawl_job(crawl_store, org_id="org-a")
        rulebook = client.post(
            f"{API_PREFIX}/deliverables/rulebooks",
            content=write_rulebook_bytes(),
            headers={"X-Org-Id": "org-b"},
        ).json()

        response = client.post(
            f"{API_PREFIX}/jobs/{job_id}/deliverable",
            json={"rulebook_id": rulebook["id"]},
            headers={"X-Org-Id": "org-a"},
        )
        assert response.status_code == 403


class TestReadingAndDownloading:
    def test_an_unknown_deliverable_is_404(self, client):
        assert (
            client.get(f"{API_PREFIX}/deliverables/nope", headers={"X-Org-Id": "org-a"}).status_code
            == 404
        )

    def test_another_orgs_deliverable_status_is_403(self, client, crawl_store):
        job_id = finished_crawl_job(crawl_store, org_id="org-a")
        accepted = client.post(
            f"{API_PREFIX}/jobs/{job_id}/deliverable", headers={"X-Org-Id": "org-a"}
        ).json()
        poll_deliverable(client, accepted["id"])

        response = client.get(
            f"{API_PREFIX}/deliverables/{accepted['id']}", headers={"X-Org-Id": "org-b"}
        )
        assert response.status_code == 403

    def test_another_orgs_deliverable_download_is_403(self, client, crawl_store):
        job_id = finished_crawl_job(crawl_store, org_id="org-a")
        accepted = client.post(
            f"{API_PREFIX}/jobs/{job_id}/deliverable", headers={"X-Org-Id": "org-a"}
        ).json()
        poll_deliverable(client, accepted["id"])

        response = client.get(
            f"{API_PREFIX}/deliverables/{accepted['id']}/download", headers={"X-Org-Id": "org-b"}
        )
        assert response.status_code == 403

    def test_downloading_before_the_build_finishes_is_409(self, client, crawl_store):
        job_id = finished_crawl_job(crawl_store)
        accepted = client.post(
            f"{API_PREFIX}/jobs/{job_id}/deliverable", headers={"X-Org-Id": "org-a"}
        ).json()

        response = client.get(
            f"{API_PREFIX}/deliverables/{accepted['id']}/download", headers={"X-Org-Id": "org-a"}
        )
        assert response.status_code in (409, 200)  # 200 only if the fast build already finished

    def test_the_list_is_org_scoped(self, client, crawl_store):
        job_id = finished_crawl_job(crawl_store, org_id="org-a")
        accepted = client.post(
            f"{API_PREFIX}/jobs/{job_id}/deliverable", headers={"X-Org-Id": "org-a"}
        ).json()
        poll_deliverable(client, accepted["id"])

        listed_a = client.get(f"{API_PREFIX}/deliverables", headers={"X-Org-Id": "org-a"}).json()
        listed_b = client.get(f"{API_PREFIX}/deliverables", headers={"X-Org-Id": "org-b"}).json()
        assert [d["id"] for d in listed_a] == [accepted["id"]]
        assert listed_b == []

    def test_a_tampered_filename_cannot_escape_the_deliverables_directory(
        self, client, crawl_store
    ):
        """Defence in depth against a tampered stored result.

        Even if a stored result named a path outside its own output
        directory, the download route refuses it rather than reading it
        (mirrors `_bundle.py`'s containment check).
        """
        job_id = finished_crawl_job(crawl_store)
        accepted = client.post(
            f"{API_PREFIX}/jobs/{job_id}/deliverable", headers={"X-Org-Id": "org-a"}
        ).json()
        poll_deliverable(client, accepted["id"])

        app_state = client.app.state.api
        app_state.deliverable_store.finish(
            accepted["id"], {"filename": "../../evil.xlsx", "site": "e.com", "pages": 1}
        )

        response = client.get(
            f"{API_PREFIX}/deliverables/{accepted['id']}/download", headers={"X-Org-Id": "org-a"}
        )
        assert response.status_code == 404


class TestBuildFromScreamingFrog:
    def test_an_empty_body_is_400(self, client):
        response = client.post(
            f"{API_PREFIX}/deliverables/from-screaming-frog",
            content=b"",
            headers={"X-Org-Id": "org-a"},
        )
        assert response.status_code == 400

    def test_a_bundle_that_fails_to_parse_ends_the_deliverable_failed_not_500(self, client):
        accepted = client.post(
            f"{API_PREFIX}/deliverables/from-screaming-frog",
            content=b"not a zip at all",
            headers={"X-Org-Id": "org-a"},
        )
        assert accepted.status_code == 202, accepted.text

        status_body = poll_deliverable(client, accepted.json()["id"])
        assert status_body["status"] == "failed"
        assert status_body["error"]

    def test_a_valid_bundle_builds_and_downloads(self, client):
        accepted = client.post(
            f"{API_PREFIX}/deliverables/from-screaming-frog",
            content=write_sf_bundle_bytes(),
            headers={"X-Org-Id": "org-a"},
        )
        assert accepted.status_code == 202, accepted.text
        deliverable_id = accepted.json()["id"]

        status_body = poll_deliverable(client, deliverable_id)
        assert status_body["status"] == "succeeded"

        download = client.get(
            f"{API_PREFIX}/deliverables/{deliverable_id}/download", headers={"X-Org-Id": "org-a"}
        )
        assert download.status_code == 200
