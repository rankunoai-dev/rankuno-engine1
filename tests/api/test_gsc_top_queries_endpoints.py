"""The top-queries fetch and the URL exports that carry it, through the HTTP API.

Search Console is replaced by a mock at the client boundary. Everything else is
real: the store, the ownership check, the sidecar, and the workbook and PDF the
download routes build. The injection and partial-data cases are asserted on the
bytes a client would actually receive.
"""

from __future__ import annotations

from io import BytesIO
from unittest.mock import MagicMock, patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from src.api import server as server_module
from src.api.server import API_PREFIX, create_app
from src.core.errors import GscAuthorizationError
from src.core.state_store import DiskJobStore, JobRecord
from src.core.url_safety import UrlSafetyPolicy
from src.integrations.gsc_schemas import GscPageQueries, GscQueryRow
from src.modules.seo.page_classifier.discovery import DiscoveryReport
from src.modules.seo.page_classifier.tool import CrawlSummary, PageClassificationOutput
from src.modules.seo.page_classifier.weights import SiteProfile, WeightProfileReport

from tests.api.conftest import TEST_SESSION_SECRET, auth_headers
from tests.modules.seo.page_classifier.report_profile_factory import (
    make_page,
    pdf_text,
)

PUBLIC_IP = "93.184.216.34"
PROPERTY = "https://e.com/"
TRIGGERS = frozenset("=+-@\t\r")


@pytest.fixture
def store(tmp_path) -> DiskJobStore:
    return DiskJobStore(tmp_path / "jobs")


@pytest.fixture
def client(store):
    app = create_app(
        store=store,
        url_policy=UrlSafetyPolicy(resolver=lambda _h: [PUBLIC_IP]),
        session_secret=TEST_SESSION_SECRET,
    )
    with TestClient(app, headers=auth_headers()) as test_client:
        yield test_client


def _finished(
    store: DiskJobStore, request: dict[str, object], org_id: str = "default"
) -> JobRecord:
    record = store.create(server_module.TOOL_NAME, request, label="e.com", org_id=org_id)
    store.finish(
        record.id,
        PageClassificationOutput(
            base_url="https://e.com/",
            site_profile=SiteProfile(),
            weight_profile=WeightProfileReport.for_site(SiteProfile()),
            discovery=DiscoveryReport(base_url="https://e.com/"),
            summary=CrawlSummary(pages_classified=2),
            pages=(make_page("https://e.com/"), make_page("https://e.com/blog/post")),
        ).model_dump(mode="json"),
    )
    return record


@pytest.fixture
def crawl(store) -> JobRecord:
    """A finished crawl whose request names a Search Console property."""
    return _finished(store, {"base_url": "https://e.com/", "gsc_property_url": PROPERTY})


def row(query: str, clicks: int = 4, page: str = "https://e.com/") -> GscQueryRow:
    return GscQueryRow(
        page=page, query=query, clicks=clicks, impressions=clicks * 10, avg_position=2.0
    )


def fetched(
    *rows: GscQueryRow, complete: bool = True, stop_reason: str = "exhausted"
) -> GscPageQueries:
    return GscPageQueries(
        property_url=PROPERTY,
        start_date="2025-10-19",
        end_date="2026-10-19",
        rows=list(rows),
        complete=complete,
        stop_reason=stop_reason,  # type: ignore[arg-type]
    )


def fake_client(result: GscPageQueries | None = None, error: Exception | None = None):
    instance = MagicMock()
    if error is not None:
        instance.fetch_top_queries.side_effect = error
    else:
        instance.fetch_top_queries.return_value = result
    return patch("src.api.server.GscApiClient", return_value=instance), instance


def fetch(client: TestClient, job_id: str):
    return client.post(f"{API_PREFIX}/jobs/{job_id}/gsc/top-queries")


def workbook(client: TestClient, job_id: str):
    response = client.get(f"{API_PREFIX}/jobs/{job_id}/urls.xlsx")
    assert response.status_code == 200
    return load_workbook(BytesIO(response.content))


class TestFetch:
    def test_the_rows_are_saved_and_only_counts_come_back(self, client, crawl, store) -> None:
        patcher, _instance = fake_client(fetched(row("shoes")))
        with patcher:
            response = fetch(client, crawl.id)

        assert response.status_code == 200
        body = response.json()
        assert body["rows"] == 1
        assert body["complete"] is True
        assert body["stop_reason"] == "exhausted"
        assert "shoes" not in response.text
        saved = store.read_gsc_queries(crawl.id)
        assert saved is not None
        assert saved["rows"][0]["query"] == "shoes"

    def test_the_property_and_the_account_come_from_the_crawl_request(self, client, store) -> None:
        record = _finished(
            store,
            {"base_url": "https://e.com/", "gsc_property_url": PROPERTY, "gsc_account": "main"},
        )
        patcher, instance = fake_client(fetched(row("shoes")))
        with patcher as constructor:
            fetch(client, record.id)

        assert constructor.call_args.kwargs == {"account": "main", "org_id": "default"}
        assert instance.fetch_top_queries.call_args.args[0] == PROPERTY

    def test_a_crawl_with_no_property_has_nothing_to_query(self, client, store) -> None:
        record = _finished(store, {"base_url": "https://e.com/"})

        response = fetch(client, record.id)

        assert response.status_code == 409
        assert "without a Search Console property" in response.json()["detail"]

    def test_an_unknown_job_is_404(self, client) -> None:
        assert fetch(client, "no-such-job").status_code == 404

    def test_another_orgs_job_is_refused(self, client, store) -> None:
        record = _finished(
            store,
            {"base_url": "https://e.com/", "gsc_property_url": PROPERTY},
            org_id="someone-else",
        )

        assert fetch(client, record.id).status_code == 403

    def test_a_property_the_account_cannot_read_is_403(self, client, crawl, store) -> None:
        patcher, _instance = fake_client(error=GscAuthorizationError("no access"))
        with patcher:
            response = fetch(client, crawl.id)

        assert response.status_code == 403
        assert store.read_gsc_queries(crawl.id) is None

    def test_a_failed_refetch_with_no_rows_keeps_the_earlier_result(
        self, client, crawl, store
    ) -> None:
        store.write_gsc_queries(crawl.id, fetched(row("kept")).model_dump(mode="json"))
        patcher, _instance = fake_client(fetched(complete=False, stop_reason="fetch_failed"))
        with patcher:
            response = fetch(client, crawl.id)

        assert response.status_code == 502
        saved = store.read_gsc_queries(crawl.id)
        assert saved is not None
        assert saved["rows"][0]["query"] == "kept"

    def test_a_partial_fetch_with_rows_is_saved_and_says_so(self, client, crawl, store) -> None:
        patcher, _instance = fake_client(
            fetched(row("shoes"), complete=False, stop_reason="fetch_failed")
        )
        with patcher:
            response = fetch(client, crawl.id)

        assert response.status_code == 200
        assert response.json()["complete"] is False
        assert store.read_gsc_queries(crawl.id) is not None

    def test_a_save_that_fails_is_503_not_a_false_success(
        self, client, crawl, store, monkeypatch
    ) -> None:
        def refuse(job_id, payload) -> None:
            raise OSError("disk full")

        monkeypatch.setattr(store, "write_gsc_queries", refuse)
        patcher, _instance = fake_client(fetched(row("shoes")))
        with patcher:
            response = fetch(client, crawl.id)

        assert response.status_code == 503


class TestExports:
    def test_top_queries_appear_in_all_urls_and_on_their_own_sheet(
        self, client, crawl, store
    ) -> None:
        store.write_gsc_queries(crawl.id, fetched(row("shoes", 4)).model_dump(mode="json"))

        book = workbook(client, crawl.id)

        assert book["All URLs"]["AA2"].value == "shoes (4)"
        assert book["Top Queries"]["C3"].value == "shoes"
        assert book["Top Queries"]["A1"].value == "GSC data complete"

    def test_a_formula_in_a_query_is_neutralised_in_the_download(
        self, client, crawl, store
    ) -> None:
        payload = '=HYPERLINK("http://evil.example","click")'
        store.write_gsc_queries(crawl.id, fetched(row(payload, 3)).model_dump(mode="json"))

        book = workbook(client, crawl.id)

        assert book["All URLs"]["AA2"].value == f"'{payload} (3)"
        assert book["Top Queries"]["C3"].value == f"'{payload}"
        for sheet in book.worksheets:
            for values in sheet.iter_rows(values_only=True):
                for value in values:
                    if isinstance(value, str):
                        assert not value or value[0] not in TRIGGERS, (sheet.title, value)

    def test_a_partial_fetch_is_announced_in_the_pdf(self, client, crawl, store) -> None:
        store.write_gsc_queries(
            crawl.id,
            fetched(row("shoes"), complete=False, stop_reason="fetch_failed").model_dump(
                mode="json"
            ),
        )

        response = client.get(f"{API_PREFIX}/jobs/{crawl.id}/urls.pdf")

        assert response.status_code == 200
        assert b"GSC data partial" in pdf_text(response.content)

    def test_a_job_with_no_fetch_exports_exactly_the_sheets_it_always_had(
        self, client, crawl
    ) -> None:
        book = workbook(client, crawl.id)

        assert book.sheetnames == ["All URLs", "By Indexability", "By HTTP Status"]
