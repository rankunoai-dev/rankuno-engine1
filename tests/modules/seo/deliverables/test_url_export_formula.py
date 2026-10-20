"""No text cell in any URL export can open with a formula trigger.

The rule is checked two ways. The first is a sweep: after writing, no string in any
sheet starts with `=`, `+`, `-` or `@` (or tab or carriage return), because every
such value must have been prefixed. The second pins where the guarded value
actually lands, so a regression that silenced the sweep by dropping a row would
still fail.
"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook
from src.integrations.gsc_schemas import GscPageQueries, GscQueryRow
from src.modules.seo.page_classifier.report_pdf import pdf_row
from src.modules.seo.page_classifier.reports import MasterURLReport
from tests.modules.seo.page_classifier.report_profile_factory import make_page, make_result

TRIGGERS = ("=", "+", "-", "@")
TRIGGER_CHARS = frozenset("=+-@\t\r")


def sweep(book) -> list[tuple[str, str]]:
    """Every string cell that opens with a trigger, as (sheet, value). Empty when safe."""
    offenders: list[tuple[str, str]] = []
    for sheet in book.worksheets:
        for values in sheet.iter_rows(values_only=True):
            for value in values:
                if isinstance(value, str) and value and value[0] in TRIGGER_CHARS:
                    offenders.append((sheet.title, value))
    return offenders


def column(book, sheet: str, index: int) -> list[object]:
    return [row[index] for row in book[sheet].iter_rows(min_row=2, values_only=True)]


def queries_with(*rows: GscQueryRow) -> GscPageQueries:
    return GscPageQueries(
        property_url="https://e.com/",
        start_date="2026-01-01",
        end_date="2026-12-31",
        rows=list(rows),
        complete=True,
        stop_reason="exhausted",
    )


@pytest.mark.parametrize("trigger", TRIGGERS)
def test_a_url_opening_with_a_trigger_is_guarded_on_each_grouped_sheet(trigger: str) -> None:
    url = f"{trigger}1+1"
    book = load_workbook(
        MasterURLReport(make_result(make_page(url), make_page("https://e.com/"))).generate()
    )

    assert sweep(book) == []
    assert f"'{url}" in column(book, "All URLs", 0)
    assert f"'{url}" in column(book, "By Indexability", 1)
    assert f"'{url}" in column(book, "By HTTP Status", 2)


@pytest.mark.parametrize("trigger", TRIGGERS)
def test_a_category_opening_with_a_trigger_is_guarded_on_all_urls(trigger: str) -> None:
    """A category is the first path segment, so a path can make it formula-shaped."""
    book = load_workbook(
        MasterURLReport(make_result(make_page(f"https://e.com/{trigger}cmd/x"))).generate()
    )

    assert sweep(book) == []
    assert f"'{trigger}cmd" in column(book, "All URLs", 27)


@pytest.mark.parametrize("trigger", TRIGGERS)
def test_a_query_opening_with_a_trigger_is_guarded_in_both_places(trigger: str) -> None:
    query = f'{trigger}HYPERLINK("http://evil.example")'
    page = make_page("https://e.com/a/")
    fetch = queries_with(
        GscQueryRow(
            page="https://e.com/a/",
            query=query,
            clicks=4,
            impressions=40,
            avg_position=2.0,
        )
    )
    book = load_workbook(MasterURLReport(make_result(page), fetch).generate())

    assert sweep(book) == []
    assert book["All URLs"]["AA2"].value == f"'{query} (4)"
    assert f"'{query}" in column(book, "Top Queries", 2)


@pytest.mark.parametrize("trigger", TRIGGERS)
def test_a_url_opening_with_a_trigger_is_guarded_in_the_pdf_too(trigger: str) -> None:
    """The PDF is not a spreadsheet, but its rows carry the same text and the same guard."""
    url = f"{trigger}1+1"

    assert pdf_row(make_page(url))[0] == f"'{url}"


def test_a_clean_export_has_nothing_to_sweep() -> None:
    book = load_workbook(MasterURLReport(make_result(make_page("https://e.com/a/"))).generate())

    assert sweep(book) == []
