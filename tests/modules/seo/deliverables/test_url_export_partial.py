"""The "GSC data partial" banner, in the workbook and in the PDF.

A partial fetch keeps every row it received, so these tests also pin that the
rows still appear: a banner over an empty export would hide a loss rather than
report it.
"""

from __future__ import annotations

from openpyxl import load_workbook
from src.integrations.gsc_schemas import GscPageQueries, GscQueryRow
from src.modules.seo.page_classifier.reports import MasterURLReport
from src.modules.seo.page_classifier.tool import GscEnrichmentReport
from tests.modules.seo.page_classifier.report_profile_factory import (
    make_page,
    make_result,
    pdf_text,
)

TOP_QUERIES_HEADER_CELL = "AA1"
"""The "Top Queries" header on All URLs: column 27, where the comment is attached."""


def partial_queries() -> GscPageQueries:
    return GscPageQueries(
        property_url="https://e.com/",
        start_date="2026-01-01",
        end_date="2026-12-31",
        rows=[
            GscQueryRow(
                page="https://e.com/a/",
                query="invoice software",
                clicks=5,
                impressions=50,
                avg_position=3.0,
            )
        ],
        complete=False,
        stop_reason="fetch_failed",
    )


def complete_queries() -> GscPageQueries:
    return partial_queries().model_copy(
        update={"complete": True, "stop_reason": "exhausted"},
    )


def test_partial_queries_are_announced_on_the_sheet_the_rows_are_on() -> None:
    report = MasterURLReport(make_result(make_page("https://e.com/a/")), partial_queries())

    book = load_workbook(report.generate())

    assert book["Top Queries"]["A1"].value.startswith("GSC data partial")
    assert "top queries (fetch_failed)" in book["Top Queries"]["A1"].value


def test_partial_queries_keep_the_rows_that_arrived() -> None:
    report = MasterURLReport(make_result(make_page("https://e.com/a/")), partial_queries())

    book = load_workbook(report.generate())

    assert book["All URLs"]["AA2"].value == "invoice software (5)"
    assert book["Top Queries"]["C3"].value == "invoice software"


def test_the_all_urls_header_carries_the_banner_as_a_comment() -> None:
    """A note row under the data would be read as a page, so the banner is a comment."""
    report = MasterURLReport(make_result(make_page("https://e.com/a/")), partial_queries())

    book = load_workbook(report.generate())

    comment = book["All URLs"][TOP_QUERIES_HEADER_CELL].comment
    assert comment is not None
    assert comment.text.startswith("GSC data partial")


def test_the_pdf_carries_the_banner_above_the_table() -> None:
    report = MasterURLReport(make_result(make_page("https://e.com/a/")), partial_queries())

    text = pdf_text(report.generate_pdf().getvalue())

    assert b"GSC data partial" in text


def test_complete_data_says_so_and_carries_no_banner() -> None:
    report = MasterURLReport(make_result(make_page("https://e.com/a/")), complete_queries())

    book = load_workbook(report.generate())

    assert book["Top Queries"]["A1"].value == "GSC data complete"
    assert book["All URLs"][TOP_QUERIES_HEADER_CELL].comment is None
    assert b"GSC data partial" not in pdf_text(report.generate_pdf().getvalue())


def test_a_partial_page_fetch_is_announced_without_any_query_sheet() -> None:
    """The page-level fetch can be partial with no query fetch at all."""
    result = make_result(make_page("https://e.com/a/")).model_copy(
        update={"gsc": GscEnrichmentReport(status="succeeded", complete=False)},
    )
    report = MasterURLReport(result)

    book = load_workbook(report.generate())

    assert book.sheetnames == ["All URLs", "By Indexability", "By HTTP Status"]
    comment = book["All URLs"][TOP_QUERIES_HEADER_CELL].comment
    assert comment is not None
    assert comment.text.startswith("GSC data partial")
    assert b"GSC data partial" in pdf_text(report.generate_pdf().getvalue())
