"""Tests for `run_deliverable_pipeline` (Step 6, ADR 0011).

Every dataset here is built in memory - `run_deliverable_pipeline` never loads
anything, so there is nothing filesystem-shaped to fixture except the
rulebook `.xlsx` files, written by `openpyxl` the same way `test_rulebook.py`
does, and the `tmp_path` output directory the real `build_workbook` writes
into. Assertions read the produced workbook back with `openpyxl` rather than
inspecting private state, so the tests exercise the same seam a caller would.
"""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

import openpyxl
import pytest
from src.modules.seo.contracts.audit import AuditDataset, AuditPage, AuditSource, Coverage
from src.modules.seo.contracts.catalogue import IssueId, Severity
from src.modules.seo.deliverables.pipeline import run_deliverable_pipeline
from src.modules.seo.deliverables.rulebook import RULEBOOK_SHEET_NAME, RulebookMissingError
from src.modules.seo.deliverables.workbook import SHEET_NOTES, SHEET_OVERVIEW, SHEET_PAGES

PRODUCED_AT = datetime(2026, 9, 9, 12, 0, tzinfo=UTC)


def identity(url: str) -> str:
    """`UrlNormalizer` stand-in for tests that don't care about normalisation."""
    return url


def full_coverage() -> dict[IssueId, Coverage]:
    return dict.fromkeys(IssueId, Coverage.NOT_MEASURED)


def make_dataset(pages: tuple[AuditPage, ...], **overrides: object) -> AuditDataset:
    base: dict[str, object] = {
        "source": AuditSource.ENGINE,
        "site": "example.com",
        "produced_at": PRODUCED_AT,
        "pages": pages,
        "issues": {},
        "coverage": full_coverage(),
    }
    base.update(overrides)
    return AuditDataset.model_validate(base)


def write_rulebook_xlsx(
    path: Path,
    *,
    rows: tuple[tuple[object, ...], ...] = (),
) -> Path:
    """A minimal rulebook workbook, one `CONTAINS` rule per row."""
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = RULEBOOK_SHEET_NAME
    sheet.append(
        ("URL Pattern", "Rule Type", "Theme 1", "Theme 2", "Language", "Business Priority")
    )
    for row in rows:
        sheet.append(row)
    book.save(path)
    return path


def read_sheet_rows(path: Path, sheet_name: str) -> list[tuple[object, ...]]:
    book = openpyxl.load_workbook(path)
    return [tuple(row) for row in book[sheet_name].iter_rows(values_only=True)]


class TestNoRulebook:
    def test_skips_theming_with_no_error_and_no_note(self, tmp_path: Path) -> None:
        dataset = make_dataset(pages=(AuditPage(url="https://example.com/"),))

        path = run_deliverable_pipeline(dataset, normalize=identity, output_dir=tmp_path)

        pages_rows = read_sheet_rows(path, SHEET_PAGES)
        assert pages_rows[1] == ("https://example.com/", None, None, None, None)
        notes_rows = read_sheet_rows(path, SHEET_NOTES)
        assert not any("rulebook" in str(cell).lower() for row in notes_rows for cell in row)


class TestRulebookApplied:
    def test_theme_fields_land_on_the_pages_sheet(self, tmp_path: Path) -> None:
        rulebook_path = write_rulebook_xlsx(
            tmp_path / "rulebook.xlsx",
            rows=(("/blog/", "CONTAINS", "Content", "Blog", "en", "High"),),
        )
        dataset = make_dataset(pages=(AuditPage(url="https://example.com/blog/post/"),))

        path = run_deliverable_pipeline(
            dataset, normalize=identity, rulebook_path=rulebook_path, output_dir=tmp_path
        )

        pages_rows = read_sheet_rows(path, SHEET_PAGES)
        assert pages_rows[1] == (
            "https://example.com/blog/post/",
            "Content",
            "Blog",
            "en",
            "High",
        )


class TestMissingRulebookStrict:
    def test_raises_rulebook_missing_error(self, tmp_path: Path) -> None:
        dataset = make_dataset(pages=(AuditPage(url="https://example.com/"),))
        missing = tmp_path / "does-not-exist.xlsx"

        with pytest.raises(RulebookMissingError):
            run_deliverable_pipeline(
                dataset, normalize=identity, rulebook_path=missing, output_dir=tmp_path
            )


class TestMissingRulebookLenient:
    def test_proceeds_and_appends_a_dataset_note(self, tmp_path: Path) -> None:
        dataset = make_dataset(pages=(AuditPage(url="https://example.com/"),))
        missing = tmp_path / "does-not-exist.xlsx"

        path = run_deliverable_pipeline(
            dataset,
            normalize=identity,
            rulebook_path=missing,
            rulebook_lenient=True,
            output_dir=tmp_path,
        )

        notes_rows = read_sheet_rows(path, SHEET_NOTES)
        assert any(
            "not found" in str(cell) and "lenient" in str(cell)
            for row in notes_rows
            for cell in row
            if cell is not None
        )


class TestWeightsAndOutputDirPassthrough:
    def test_custom_weights_change_the_overview_penalty_total(self, tmp_path: Path) -> None:
        pages = (AuditPage(url="https://example.com/a/"),)
        coverage = full_coverage()
        coverage[IssueId.H1_MISSING] = Coverage.MEASURED
        dataset = make_dataset(
            pages=pages,
            issues={IssueId.H1_MISSING: frozenset({"https://example.com/a/"})},
            coverage=coverage,
        )

        default_out = tmp_path / "default"
        custom_out = tmp_path / "custom"
        default_path = run_deliverable_pipeline(dataset, normalize=identity, output_dir=default_out)
        custom_path = run_deliverable_pipeline(
            dataset,
            normalize=identity,
            weights={Severity.ISSUE: 99, Severity.WARNING: 1, Severity.OPPORTUNITY: 1},
            output_dir=custom_out,
        )

        assert default_path.parent == default_out
        assert custom_path.parent == custom_out

        default_overview = read_sheet_rows(default_path, SHEET_OVERVIEW)
        custom_overview = read_sheet_rows(custom_path, SHEET_OVERVIEW)
        # H1 is Category.H1, Severity.ISSUE: default weight 3 vs custom weight 99.
        default_total = next(row[1] for row in default_overview if row[0] == "H1")
        custom_total = next(row[1] for row in custom_overview if row[0] == "H1")
        assert default_total == 3
        assert custom_total == 99
