"""End-to-end masterfile build from a real Screaming Frog export bundle (0107).

Every other masterfile test constructs a service against a directory that a
test just wrote. Nothing proved a masterfile could be built from the only
Screaming Frog export this platform actually stores: one zip, encrypted at
rest, never extracted. Before this cycle it could not be - `run_masterfile`
took a filesystem path and the services joined it per CSV - so every test in
this module would have failed.

The member names here are asserted to be a subset of
`ALLOWED_BUNDLE_FILENAMES`, which is derived from `ISSUE_CATALOGUE`'s
`sf_sources` filenames plus the spine - deliberately *not* from the Screaming
Frog CLI argument strings, a derivation build-log 0108 rejected because for
seven files those arguments cannot produce the filename Screaming Frog writes.
That is what makes this bundle *real* rather than merely zip-shaped: a name not
in that set could never arrive through `POST /workers/jobs/{id}/upload`.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import openpyxl
from src.core.state_store import DiskJobStore, JobStatus
from src.modules.seo.deliverables.build_runner import run_masterfile
from src.modules.seo.deliverables.masterfile_source import (
    source_for_bundle_bytes,
    source_for_path,
)
from src.modules.seo.screaming_frog_control.upload_manifest import ALLOWED_BUNDLE_FILENAMES

URL = "https://example.com/pricing"

EXPORT: dict[str, str] = {
    "internal_all.csv": (
        "Address,Content Type,Status Code,Indexability,Indexability Status,Inlinks\n"
        f"{URL},text/html,200,Indexable,,7\n"
    ),
    "h1_missing.csv": f"Address,Occurrences\n{URL},0\n",
}


def bundle_bytes(members: dict[str, str] | None = None) -> bytes:
    """A zip shaped exactly like a worker upload: BOM-prefixed UTF-8 CSVs."""
    chosen = EXPORT if members is None else members
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, body in chosen.items():
            archive.writestr(name, body.encode("utf-8-sig"))
    return buffer.getvalue()


def test_the_fixture_is_a_real_screaming_frog_export() -> None:
    assert set(EXPORT) <= ALLOWED_BUNDLE_FILENAMES


def read_rows(path: Path) -> list[tuple[object, ...]]:
    book = openpyxl.load_workbook(path)
    sheet = book.active
    assert sheet is not None
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


def test_a_masterfile_builds_from_a_bundle_held_only_in_memory(tmp_path: Path) -> None:
    store = DiskJobStore(tmp_path / "deliverables")
    record = store.create("seo.deliverables.workbook", {"service_slug": "h1"}, org_id="org-a")
    payload = bundle_bytes()

    run_masterfile(
        store,
        record.id,
        "sf-job-1",
        "h1",
        lambda: source_for_bundle_bytes(payload, "sf-job-1"),
        None,
    )

    finished = store.get(record.id)
    assert finished.status is JobStatus.SUCCEEDED, finished.error
    result = store.read_result(record.id)
    assert result["filename"] == "h1.xlsx"
    assert result["source_job_id"] == "sf-job-1"

    workbook = store.root / record.id / "h1.xlsx"
    rows = read_rows(workbook)
    assert ("Total Affected Pages", 1) in [row[:2] for row in rows]
    assert any(row[0] == URL for row in rows), rows


def test_the_bundle_is_never_written_to_disk(tmp_path: Path) -> None:
    """The reason this path exists at all: bundles are encrypted at rest."""
    store = DiskJobStore(tmp_path / "deliverables")
    record = store.create("seo.deliverables.workbook", {"service_slug": "h1"}, org_id="org-a")
    payload = bundle_bytes()

    run_masterfile(
        store,
        record.id,
        "sf-job-1",
        "h1",
        lambda: source_for_bundle_bytes(payload, "sf-job-1"),
        None,
    )

    written = {path.suffix for path in tmp_path.rglob("*") if path.is_file()}
    assert written == {".json", ".xlsx"}, written


def test_a_dynamic_export_is_discovered_from_the_bundle(tmp_path: Path) -> None:
    """`custom_extraction_*` has one file per extractor; no glob in a zip."""
    store = DiskJobStore(tmp_path / "deliverables")
    record = store.create(
        "seo.deliverables.workbook", {"service_slug": "custom_extraction"}, org_id="org-a"
    )
    members = dict(EXPORT)
    members["custom_extraction_authors.csv"] = f"Address,Author 1\n{URL},Ada\n"
    payload = bundle_bytes(members)

    run_masterfile(
        store,
        record.id,
        "sf-job-1",
        "custom_extraction",
        lambda: source_for_bundle_bytes(payload, "sf-job-1"),
        None,
    )

    finished = store.get(record.id)
    assert finished.status is JobStatus.SUCCEEDED, finished.error
    book = openpyxl.load_workbook(store.root / record.id / "custom_extraction.xlsx")
    assert book.sheetnames == ["authors"]


def test_an_export_directory_still_builds(tmp_path: Path) -> None:
    """The directory source stays supported: `build_deliverable.py` uses it."""
    export = tmp_path / "sf_export"
    export.mkdir()
    for name, body in EXPORT.items():
        (export / name).write_text(body, encoding="utf-8")
    store = DiskJobStore(tmp_path / "deliverables")
    record = store.create("seo.deliverables.workbook", {"service_slug": "h1"}, org_id="org-a")

    run_masterfile(store, record.id, "job-1", "h1", lambda: source_for_path(export), None)

    finished = store.get(record.id)
    assert finished.status is JobStatus.SUCCEEDED, finished.error
    assert any(row[0] == URL for row in read_rows(store.root / record.id / "h1.xlsx"))


def test_a_corrupt_bundle_fails_the_build_rather_than_crashing(tmp_path: Path) -> None:
    store = DiskJobStore(tmp_path / "deliverables")
    record = store.create("seo.deliverables.workbook", {"service_slug": "h1"}, org_id="org-a")

    run_masterfile(
        store,
        record.id,
        "sf-job-1",
        "h1",
        lambda: source_for_bundle_bytes(b"not a zip", "sf-job-1"),
        None,
    )

    finished = store.get(record.id)
    assert finished.status is JobStatus.FAILED
    assert finished.error is not None
    assert "not-a-zip" in finished.error


def test_an_unknown_service_fails_the_build(tmp_path: Path) -> None:
    store = DiskJobStore(tmp_path / "deliverables")
    record = store.create("seo.deliverables.workbook", {"service_slug": "nope"}, org_id="org-a")
    payload = bundle_bytes()

    run_masterfile(
        store, record.id, "sf-job-1", "nope", lambda: source_for_bundle_bytes(payload, "x"), None
    )

    finished = store.get(record.id)
    assert finished.status is JobStatus.FAILED
    assert finished.error is not None
    assert "nope" in finished.error
