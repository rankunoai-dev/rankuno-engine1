"""The top-queries cell, the partial banner, and the Top Queries sheet layout."""

from __future__ import annotations

import pytest
from openpyxl import Workbook
from src.integrations.gsc_schemas import GscPageQueries, GscQueryRow
from src.modules.seo.performance.query_sheet import (
    GSC_COMPLETE_NOTE,
    partial_banner,
    top_queries_cell,
    write_top_queries_sheet,
)


def row(query: str, clicks: int = 1, page: str = "https://e.com/a/") -> GscQueryRow:
    return GscQueryRow(
        page=page,
        query=query,
        clicks=clicks,
        impressions=clicks * 10,
        avg_position=2.0,
    )


def fetch(*, complete: bool = True, stop_reason: str = "exhausted") -> GscPageQueries:
    return GscPageQueries(
        property_url="https://e.com/",
        start_date="2026-01-01",
        end_date="2026-12-31",
        rows=[],
        complete=complete,
        stop_reason=stop_reason,  # type: ignore[arg-type]
    )


class TestCell:
    def test_query_and_clicks_are_joined_in_rank_order(self) -> None:
        assert (
            top_queries_cell([row("red shoes", 12), row("blue", 3)]) == "red shoes (12); blue (3)"
        )

    def test_no_queries_is_a_blank_cell_not_a_zero(self) -> None:
        assert top_queries_cell([]) is None

    @pytest.mark.parametrize("trigger", ["=", "+", "-", "@"])
    def test_a_query_opening_with_a_trigger_is_guarded(self, trigger: str) -> None:
        text = f'{trigger}HYPERLINK("http://x")'

        assert top_queries_cell([row(text, 4)]) == f"'{text} (4)"


class TestBanner:
    def test_complete_data_has_no_banner(self) -> None:
        assert partial_banner(page_metrics_complete=True, queries=fetch()) is None

    def test_partial_page_metrics_are_named(self) -> None:
        text = partial_banner(page_metrics_complete=False, queries=None)

        assert text is not None
        assert text.startswith("GSC data partial")
        assert "page metrics" in text

    def test_partial_queries_name_their_stop_reason(self) -> None:
        text = partial_banner(
            page_metrics_complete=True,
            queries=fetch(complete=False, stop_reason="fetch_failed"),
        )

        assert text is not None
        assert "top queries (fetch_failed)" in text


class TestSheet:
    def test_status_line_then_header_then_one_row_per_rank(self) -> None:
        book = Workbook()
        ws = book.active
        assert ws is not None
        rows = {"https://e.com/a/": [row("x", 5), row("y", 2)]}

        write_top_queries_sheet(ws, ["https://e.com/a/"], rows, None)

        values = list(ws.iter_rows(values_only=True))
        assert values[0][0] == GSC_COMPLETE_NOTE
        assert values[1][:6] == ("URL", "Rank", "Query", "Clicks", "Impressions", "Position")
        assert values[2][:3] == ("https://e.com/a/", 1, "x")
        assert values[3][:3] == ("https://e.com/a/", 2, "y")

    def test_a_banner_takes_the_status_line(self) -> None:
        book = Workbook()
        ws = book.active
        assert ws is not None

        write_top_queries_sheet(ws, [], {}, "GSC data partial: example")

        assert ws["A1"].value == "GSC data partial: example"

    @pytest.mark.parametrize("trigger", ["=", "+", "-", "@"])
    def test_a_url_and_a_query_opening_with_a_trigger_are_both_guarded(self, trigger: str) -> None:
        book = Workbook()
        ws = book.active
        assert ws is not None
        url = f"{trigger}evil"
        query = f"{trigger}cmd"

        write_top_queries_sheet(ws, [url], {url: [row(query, 1, page=url)]}, None)

        assert ws["A3"].value == f"'{url}"
        assert ws["C3"].value == f"'{query}"
