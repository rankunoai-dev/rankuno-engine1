"""HTTP status capture: the URL-list exports must say what each fetch returned.

Before this, `SiteGraph.record_fetch` read `FetchResult.status_code` once, fed it
to `indexability_of`, and dropped it, so every row of the XLSX/PDF "HTTP Status"
column read "Unknown". These tests crawl a small stub site through the governed
tool on both crawl paths and check the exports, end to end, against what the
stub actually answered.
"""

from __future__ import annotations

import base64
import re
import zlib
from pathlib import Path
from typing import Any

import httpx
import pytest
from openpyxl import load_workbook
from src.api.server import _placeholder_profile
from src.core.url_safety import UrlSafetyPolicy
from src.modules.seo.page_classifier.discovery import (
    OUTCOME_GUARDRAIL_REFUSED,
    OUTCOME_NOT_FETCHED,
    OUTCOME_ROBOTS_DISALLOWED,
    OUTCOME_TRANSPORT_REFUSED,
    OUTCOME_TRANSPORT_TIMEOUT,
    DiscoveredNode,
    discover_site,
    record_fetch_exception,
)
from src.modules.seo.page_classifier.fetch_status import (
    LABEL_NO_RESPONSE,
    LABEL_NOT_FETCHED,
    LABEL_NOT_RECORDED,
    first_redirect_label,
    http_status_label,
)
from src.modules.seo.page_classifier.reports import MasterURLReport
from src.modules.seo.page_classifier.schemas import FullPageIntelligenceProfile
from src.modules.seo.page_classifier.tool import (
    PageClassificationInput,
    PageClassificationOutput,
    PageClassificationTool,
)

from tests.modules.seo.golden_site_factory import (
    BASE,
    GoldenFetcher,
    build_fetcher,
    build_site,
    golden_settings,
)
from tests.modules.seo.test_golden_extract_at_fetch import STATUS_FIELDS

PUBLIC_IP = "93.184.216.34"
ROOT = "https://e.com"
HOME = (
    "<html><head><title>Home</title></head><body>"
    '<a href="/ok">ok</a><a href="/moved">moved</a><a href="/gone">gone</a>'
    '<a href="/busted">busted</a><a href="/refused-conn">refused</a>'
    '<a href="/private/secret">private</a><a href="/bait">bait</a>'
    '<a href="/download">download</a>'
    "</body></html>"
)
SITEMAP = (
    '<?xml version="1.0"?><urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">'
    f"<url><loc>{ROOT}/ok</loc></url><url><loc>{ROOT}/only-in-sitemap</loc></url></urlset>"
)

# path -> (final status, first redirect, outcome) the stub must produce.
EXPECTED: dict[str, tuple[int | None, int | None, str]] = {
    "/ok": (200, None, "ok"),
    "/moved": (200, 301, "ok"),
    "/gone": (404, None, "not_found"),
    "/busted": (503, None, "server_error"),
    "/download": (200, None, "not_html"),
    "/refused-conn": (None, None, OUTCOME_TRANSPORT_REFUSED),
    "/private/secret": (None, None, OUTCOME_ROBOTS_DISALLOWED),
    "/bait": (None, None, OUTCOME_GUARDRAIL_REFUSED),
    "/only-in-sitemap": (None, None, OUTCOME_NOT_FETCHED),
}
LABELS = {
    "/refused-conn": LABEL_NO_RESPONSE,
    "/bait": LABEL_NO_RESPONSE,
    "/private/secret": LABEL_NOT_FETCHED,
    "/only-in-sitemap": LABEL_NOT_FETCHED,
}


def _handler(request: httpx.Request) -> httpx.Response:
    path = request.url.path
    html = {"content-type": "text/html"}
    if path == "/robots.txt":
        return httpx.Response(
            200, text=f"User-agent: *\nDisallow: /private/\nSitemap: {ROOT}/sitemap.xml\n"
        )
    if path == "/sitemap.xml":
        return httpx.Response(200, text=SITEMAP, headers={"content-type": "application/xml"})
    if path == "/":
        return httpx.Response(200, text=HOME, headers=html)
    if path == "/ok":
        return httpx.Response(200, text="<html><body>ok</body></html>", headers=html)
    if path == "/moved":
        return httpx.Response(301, headers={"location": f"{ROOT}/moved-to"})
    if path == "/moved-to":
        return httpx.Response(200, text="<html><body>landed</body></html>", headers=html)
    if path == "/busted":
        return httpx.Response(503, text="down", headers=html)
    if path == "/download":
        return httpx.Response(200, content=b"%PDF", headers={"content-type": "application/pdf"})
    if path == "/bait":
        return httpx.Response(302, headers={"location": "http://169.254.169.254/"})
    if path == "/refused-conn":
        raise httpx.ConnectError("refused", request=request)
    if path == "/slow":
        return httpx.Response(200, text="<html><body>slow</body></html>", headers=html)
    return httpx.Response(404, text="gone", headers=html)


def _fetcher(tmp: Path) -> GoldenFetcher:
    return GoldenFetcher(
        settings=golden_settings(tmp),
        url_policy=UrlSafetyPolicy(resolver=lambda host: [PUBLIC_IP]),
        transport=httpx.MockTransport(_handler),
        async_transport=httpx.MockTransport(_handler),
    )


def _crawl(tmp: Path, *, use_async: bool) -> PageClassificationOutput:
    tool = PageClassificationTool(fetcher=_fetcher(tmp))
    result = tool.run(PageClassificationInput(base_url=ROOT, use_async_crawl=use_async))
    assert result.ok, result.error
    assert isinstance(result.data, PageClassificationOutput)
    return result.data


@pytest.fixture(scope="module", params=[False, True], ids=["serial", "async"])
def crawled(request: pytest.FixtureRequest, tmp_path_factory: pytest.TempPathFactory) -> Any:
    return _crawl(tmp_path_factory.mktemp("status"), use_async=request.param)


def _by_path(output: PageClassificationOutput) -> dict[str, FullPageIntelligenceProfile]:
    return {page.url.removeprefix(ROOT): page for page in output.pages}


def _pdf_text(pdf: bytes) -> bytes:
    chunks: list[bytes] = []
    for match in re.finditer(rb"stream\r?\n(.*?)endstream", pdf, re.DOTALL):
        raw = match.group(1).rstrip(b"\r\n")
        if raw.endswith(b"~>"):
            raw = raw[:-2]
        try:
            chunks.append(zlib.decompress(base64.a85decode(raw, adobe=False)))
        except (ValueError, zlib.error):
            continue
    return b"".join(chunks)


class TestCapture:
    def test_each_page_carries_the_status_the_stub_answered(self, crawled: Any) -> None:
        pages = _by_path(crawled)
        for path, (status, first, outcome) in EXPECTED.items():
            page = pages[path]
            assert (page.http_status, page.first_redirect_status, page.fetch_outcome) == (
                status,
                first,
                outcome,
            ), path

    def test_a_redirect_is_recorded_not_applied(self, crawled: Any) -> None:
        page = _by_path(crawled)["/moved"]
        assert page.url == f"{ROOT}/moved"
        assert page.final_url == f"{ROOT}/moved-to"
        assert page.redirect_chain == (f"{ROOT}/moved",)

    def test_xlsx_all_urls_shows_the_right_cell_for_every_page(self, crawled: Any) -> None:
        book = load_workbook(MasterURLReport(crawled).generate())
        rows = list(book["All URLs"].iter_rows(values_only=True))
        header = list(rows[0])
        status_col, first_col = header.index("HTTP Status"), header.index("First Redirect")
        by_url = {row[0].removeprefix(ROOT): row for row in rows[1:]}
        for path, (status, first, _) in EXPECTED.items():
            assert by_url[path][status_col] == (str(status) if status else LABELS[path]), path
            assert by_url[path][first_col] == (str(first) if first else "—"), path

    def test_pdf_shows_the_same_cells_and_the_legend(self, crawled: Any) -> None:
        report = MasterURLReport(crawled)
        text = _pdf_text(report.generate_pdf().getvalue())
        for label in (LABEL_NO_RESPONSE, LABEL_NOT_FETCHED, "503", "404", "301"):
            assert label.encode() in text, label
        assert b"First Redirect" in text and b"last attempt" in text
        for page in crawled.pages:
            assert report._pdf_row(page)[3] == http_status_label(page)
            assert report._pdf_row(page)[4] == first_redirect_label(page)

    def test_by_http_status_sheet_groups_codes_then_labels(self, crawled: Any) -> None:
        book = load_workbook(MasterURLReport(crawled).generate())
        groups = [row[0] for row in book["By HTTP Status"].iter_rows(min_row=2, values_only=True)]
        ordered = list(dict.fromkeys(groups))
        numeric = [g for g in ordered if g.isdigit()]
        assert numeric == sorted(numeric, key=int)
        assert ordered[: len(numeric)] == numeric
        assert ordered[len(numeric) :] == [LABEL_NOT_FETCHED, LABEL_NO_RESPONSE]


def test_serial_and_async_record_identical_status_and_outcome(tmp_path: Path) -> None:
    serial = _by_path(_crawl(tmp_path / "s", use_async=False))
    parallel = _by_path(_crawl(tmp_path / "a", use_async=True))
    assert serial.keys() == parallel.keys()
    for path in serial:
        left, right = serial[path], parallel[path]
        assert (left.http_status, left.first_redirect_status, left.fetch_outcome) == (
            right.http_status,
            right.first_redirect_status,
            right.fetch_outcome,
        ), path


class TestLegacyResults:
    def test_a_stored_result_without_the_fields_loads_and_renders_the_legacy_label(
        self, tmp_path: Path
    ) -> None:
        output = _crawl(tmp_path, use_async=False)
        dumped = output.model_dump(mode="json")
        for page in dumped["pages"]:
            for field in STATUS_FIELDS:
                del page[field]
        old = PageClassificationOutput.model_validate(dumped)
        assert all(page.http_status is None and page.fetch_outcome is None for page in old.pages)
        book = load_workbook(MasterURLReport(old).generate())
        statuses = {row[3] for row in book["All URLs"].iter_rows(min_row=2, values_only=True)}
        assert statuses == {LABEL_NOT_RECORDED}
        assert b"Not recorded" in _pdf_text(MasterURLReport(old).generate_pdf().getvalue())

    def test_a_checkpoint_placeholder_profile_renders_the_legacy_label(self) -> None:
        page = _placeholder_profile(f"{ROOT}/x", None)
        assert http_status_label(page) == LABEL_NOT_RECORDED
        assert first_redirect_label(page) == "—"

    def test_the_new_fields_default_to_none_on_every_model(self) -> None:
        node = DiscoveredNode(url=f"{ROOT}/x", normalized="/x")
        assert (node.http_status, node.first_redirect_status, node.fetch_outcome) == (
            None,
            None,
            None,
        )


class TestRecordSemantics:
    def test_indexability_is_unchanged_by_the_capture(self, tmp_path: Path) -> None:
        """The verdict still comes from `indexability_of` with the same inputs."""
        from src.modules.seo.page_classifier.signal_parsers import (
            extract_robots_directives,
            indexability_of,
        )

        fetcher = _fetcher(tmp_path)
        graph, _ = discover_site(fetcher, ROOT)
        for url in (f"{ROOT}/gone", f"{ROOT}/busted", f"{ROOT}/moved", f"{ROOT}/ok"):
            result = fetcher.fetch(url)
            node = next(n for n in graph._nodes.values() if n.url == url)
            expected = indexability_of(
                url,
                status_code=result.status_code,
                directives=extract_robots_directives(result.body, result.headers),
                canonical_url=node.canonical_url,
                is_html=result.is_html,
                redirected_to=result.final_url,
            )
            assert (node.indexability, node.indexability_reason) == expected

    def test_a_failure_after_a_success_clears_the_stale_status(self, tmp_path: Path) -> None:
        """Last attempt wins: a later failed attempt must not leave an old code."""
        fetcher = _fetcher(tmp_path)
        graph, _ = discover_site(fetcher, ROOT)
        node = next(n for n in graph._nodes.values() if n.url == f"{ROOT}/moved")
        assert node.http_status == 200 and node.first_redirect_status == 301
        graph.record_attempt_failure(node.url, OUTCOME_TRANSPORT_REFUSED)
        assert (node.http_status, node.first_redirect_status) == (None, None)
        assert node.fetch_outcome == OUTCOME_TRANSPORT_REFUSED

    def test_a_timeout_is_no_response_on_either_path_setting(self, tmp_path: Path) -> None:
        graph, _ = discover_site(_fetcher(tmp_path), ROOT)
        for classify in (False, True):
            url = f"{ROOT}/ok"
            record_fetch_exception(
                graph,
                url,
                httpx.ReadTimeout("tarpit"),
                classify_transport=classify,
                mark_node=True,
            )
            node = next(n for n in graph._nodes.values() if n.url == url)
            assert node.fetch_outcome == OUTCOME_TRANSPORT_TIMEOUT
            assert node.http_status is None

    def test_a_sitemap_or_cms_endpoint_failure_does_not_mark_a_page(self, tmp_path: Path) -> None:
        graph, _ = discover_site(_fetcher(tmp_path), ROOT)
        before = next(n for n in graph._nodes.values() if n.url == f"{ROOT}/ok").http_status
        record_fetch_exception(graph, f"{ROOT}/ok", httpx.ReadTimeout("x"), classify_transport=True)
        after = next(n for n in graph._nodes.values() if n.url == f"{ROOT}/ok")
        assert after.http_status == before == 200

    def test_status_ints_are_shared_objects(self, tmp_path: Path) -> None:
        graph, _ = discover_site(_fetcher(tmp_path), ROOT)
        gone = [n.http_status for n in graph._nodes.values() if n.http_status == 404]
        assert gone
        assert all(code is gone[0] for code in gone)

    def test_evicted_loop_nodes_never_reach_the_profiles(self, tmp_path: Path) -> None:
        """An evicted node was fetched and carried a status; it must vanish whole."""
        requested: list[str] = []

        class Recording(dict):  # type: ignore[type-arg]
            def get(self, key: Any, default: Any = None) -> Any:
                requested.append(key)
                return super().get(key, default)

        site = Recording(build_site())
        graph, _ = discover_site(build_fetcher(golden_settings(tmp_path), site), BASE)
        kept = {node.url.removeprefix(BASE) for node in graph._nodes.values()}
        evicted = {path for path in requested if "tools/roi/calc" in path} - kept
        assert evicted, "precondition: the golden site's loop evicts fetched nodes"
        evidence_paths = {e.url.removeprefix(BASE) for e in graph.to_page_evidence()}
        assert not evicted & evidence_paths
