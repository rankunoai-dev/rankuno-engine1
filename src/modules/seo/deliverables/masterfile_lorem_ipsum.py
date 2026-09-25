"""Lorem Ipsum masterfile service (placeholder/dummy text detection).

Reads lorem_ipsum.csv file and generates a single-sheet XLSX showing:
1. Summary: Placeholder text counts
2. Detailed Data: One row per affected URL, sorted by impressions

Source CSV:
- lorem_ipsum.csv
"""

from __future__ import annotations

import io
from typing import Any, Final

import pandas as pd  # type: ignore[import-untyped]
from openpyxl import Workbook

from src.core.logger import get_logger
from src.modules.seo.deliverables.masterfile_base import (
    MasterfileMetadata,
    MasterfileService,
    gc,
    read_csv_safe,
    safe_cell,
)

__all__ = ["LoremIpsumService"]

_logger = get_logger(__name__)

_LOREM_IPSUM_FILE: Final[str] = "lorem_ipsum.csv"


class LoremIpsumService(MasterfileService):
    """Generate lorem ipsum masterfile."""

    @property
    def metadata(self) -> MasterfileMetadata:
        return MasterfileMetadata(
            slug="lorem_ipsum",
            label="Lorem Ipsum",
            sheets=1,
            is_complex=False,
        )

    def _read_lorem_ipsum(self) -> pd.DataFrame | None:
        """Read lorem ipsum CSV."""
        return read_csv_safe(self.sf_export_dir / _LOREM_IPSUM_FILE)

    def generate(self) -> bytes:
        """Generate lorem ipsum XLSX."""
        # Read data
        lorem_df = self._read_lorem_ipsum()
        internal_map = self._build_internal_map()
        gsc_map = self._build_gsc_map()

        if lorem_df is None or lorem_df.empty:
            wb = Workbook()
            ws = wb.active
            if ws:
                ws.title = "Lorem Ipsum"
                ws.append(["No placeholder text found"])
        else:
            try:
                address_idx = gc(lorem_df.columns.tolist(), "Address")
            except KeyError:
                wb = Workbook()
                ws = wb.active
                if ws:
                    ws.title = "Lorem Ipsum"
                    ws.append(["Error reading lorem ipsum CSV"])
                output = io.BytesIO()
                wb.save(output)
                return output.getvalue()

            urls_with_lorem = []
            for _, row in lorem_df.iterrows():
                url = str(row.iloc[address_idx])

                # Enrich
                internal_data = internal_map.get(url, {}) if internal_map else {}
                gsc_data = gsc_map.get(url, {}) if gsc_map else {}

                # Filter: indexable=True
                indexability = internal_data.get("indexability")
                if indexability != "Indexable":
                    continue

                urls_with_lorem.append(
                    {
                        "url": url,
                        "indexability": indexability,
                        "inlinks": internal_data.get("inlinks"),
                        "impressions": gsc_data.get("impressions", 0),
                        "clicks": gsc_data.get("clicks", 0),
                    }
                )

            # Build workbook
            wb = Workbook()
            ws = wb.active
            if ws:
                ws.title = "Lorem Ipsum"
                self._write_lorem_ipsum_sheet(ws, urls_with_lorem)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    def _write_lorem_ipsum_sheet(
        self, ws: object, urls_with_lorem: list[dict[str, Any]]
    ) -> None:
        """Write lorem ipsum sheet."""
        ws.append([])  # type: ignore[attr-defined]
        ws.append(["LOREM IPSUM / PLACEHOLDER TEXT"])  # type: ignore[attr-defined]
        ws.append([])  # type: ignore[attr-defined]

        # Summary section
        ws.append(["Summary"])  # type: ignore[attr-defined]
        ws.append(["Total Affected Pages", len(urls_with_lorem)])  # type: ignore[attr-defined]
        ws.append([])  # type: ignore[attr-defined]

        # Detailed Data section
        ws.append(["Detailed Data"])  # type: ignore[attr-defined]
        ws.append(["URL", "Inlinks", "Impressions", "Clicks"])  # type: ignore[attr-defined]

        # Sort by impressions descending
        sorted_urls = sorted(urls_with_lorem, key=lambda x: x.get("impressions", 0), reverse=True)
        for url_data in sorted_urls:
            ws.append(  # type: ignore[attr-defined]
                [
                    safe_cell(url_data["url"]),
                    url_data.get("inlinks"),
                    url_data.get("impressions", 0),
                    url_data.get("clicks", 0),
                ]
            )
