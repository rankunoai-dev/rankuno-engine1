"""Pagination masterfile service (rel=next/prev tag analysis).

Reads its catalogue sources and generates a single-sheet XLSX showing:
1. Summary: Pagination issue counts
2. Detailed Data: One row per affected URL, sorted by impressions

Source CSVs: `SOURCE_FILES`, derived from IssueCategory.PAGINATION through
`contracts/sources.py`. Never written out here - a hand-kept second
copy of that list is what build-log 0116 found wrong in this file.
"""

from __future__ import annotations

import io
from typing import Any

import pandas as pd  # type: ignore[import-untyped]
from openpyxl import Workbook

from src.core.logger import get_logger
from src.modules.seo.contracts.issue_ids import IssueCategory
from src.modules.seo.contracts.sources import sources_for_categories
from src.modules.seo.deliverables.masterfile_base import (
    MasterfileMetadata,
    MasterfileService,
    gc,
    safe_cell,
)

__all__ = ["PaginationService"]

_logger = get_logger(__name__)


class PaginationService(MasterfileService):
    """Generate pagination masterfile."""

    SOURCE_FILES = sources_for_categories(IssueCategory.PAGINATION)

    @property
    def metadata(self) -> MasterfileMetadata:
        return MasterfileMetadata(
            slug="pagination",
            label="Pagination",
            sheets=1,
            is_complex=False,
        )

    def _read_all_pagination(self) -> pd.DataFrame | None:
        """Read and combine all pagination CSVs."""
        return self._read_issue_frames()

    def generate(self) -> bytes:
        """Generate pagination XLSX."""
        # Read data
        pagination_df = self._read_all_pagination()
        internal_map = self._build_internal_map()
        gsc_map = self._build_gsc_map()

        if pagination_df is None or pagination_df.empty:
            wb = Workbook()
            ws = wb.active
            if ws:
                ws.title = "Pagination"
                ws.append(["No pagination data found"])
        else:
            try:
                address_idx = gc(pagination_df.columns.tolist(), "Address")
            except KeyError:
                wb = Workbook()
                ws = wb.active
                if ws:
                    ws.title = "Pagination"
                    ws.append(["Error reading pagination CSV"])
                output = io.BytesIO()
                wb.save(output)
                return output.getvalue()

            urls_with_pagination = []
            for _, row in pagination_df.iterrows():
                url = str(row.iloc[address_idx])

                # Enrich
                internal_data = internal_map.get(url, {}) if internal_map else {}
                gsc_data = self._gsc_lookup(gsc_map, url)

                indexability = internal_data.get("indexability")
                if not self._is_reportable(indexability):
                    continue

                urls_with_pagination.append(
                    {
                        "url": url,
                        "indexability": indexability,
                        "inlinks": internal_data.get("inlinks"),
                        "impressions": gsc_data.get("impressions", 0),
                        "clicks": gsc_data.get("clicks", 0),
                    }
                )

            # Build workbook
            wb = Workbook()
            ws = wb.active
            if ws:
                ws.title = "Pagination"
                self._write_pagination_sheet(ws, urls_with_pagination)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    def _write_pagination_sheet(
        self, ws: object, urls_with_pagination: list[dict[str, Any]]
    ) -> None:
        """Write pagination sheet."""
        ws.append([])  # type: ignore[attr-defined]
        ws.append(["PAGINATION"])  # type: ignore[attr-defined]
        ws.append([])  # type: ignore[attr-defined]

        # Summary section
        ws.append(["Summary"])  # type: ignore[attr-defined]
        ws.append(["Total Affected Pages", len(urls_with_pagination)])  # type: ignore[attr-defined]
        ws.append([])  # type: ignore[attr-defined]

        # Detailed Data section
        ws.append(["Detailed Data"])  # type: ignore[attr-defined]
        ws.append(["URL", "Inlinks", "Impressions", "Clicks"])  # type: ignore[attr-defined]

        # Sort by impressions descending
        sorted_urls = sorted(
            urls_with_pagination, key=lambda x: x.get("impressions", 0), reverse=True
        )
        for url_data in sorted_urls:
            ws.append(  # type: ignore[attr-defined]
                [
                    safe_cell(url_data["url"]),
                    url_data.get("inlinks"),
                    url_data.get("impressions", 0),
                    url_data.get("clicks", 0),
                ]
            )
