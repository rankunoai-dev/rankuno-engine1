"""Non-Functional Internal Links masterfile service (broken internal links).

Reads non_functional_internal_links.csv file and generates a single-sheet XLSX showing:
1. Summary: Broken link counts
2. Detailed Data: One row per affected URL, sorted by impressions

Source CSV:
- non_functional_internal_links.csv
"""

from __future__ import annotations

import io
from typing import Any, Final

import pandas as pd  # type: ignore[import-untyped]
from openpyxl import Workbook

from src.core.logger import get_logger
from src.modules.seo.deliverables.masterfile_base import (
    MasterfileMetadata,
    MasterfileService,
    gc,
    read_csv_safe,
    safe_cell,
)

__all__ = ["NonFunctionalInternalLinksService"]

_logger = get_logger(__name__)

_NON_FUNCTIONAL_LINKS_FILE: Final[str] = "non_functional_internal_links.csv"


class NonFunctionalInternalLinksService(MasterfileService):
    """Generate non-functional internal links masterfile."""

    @property
    def metadata(self) -> MasterfileMetadata:
        return MasterfileMetadata(
            slug="non_functional_internal_links",
            label="Non-Functional Internal Links",
            sheets=1,
            is_complex=False,
        )

    def _read_non_functional_links(self) -> pd.DataFrame | None:
        """Read non-functional internal links CSV."""
        return read_csv_safe(self.sf_export_dir / _NON_FUNCTIONAL_LINKS_FILE)

    def generate(self) -> bytes:
        """Generate non-functional internal links XLSX."""
        # Read data
        links_df = self._read_non_functional_links()
        internal_map = self._build_internal_map()
        gsc_map = self._build_gsc_map()

        if links_df is None or links_df.empty:
            wb = Workbook()
            ws = wb.active
            if ws:
                ws.title = "Broken Links"
                ws.append(["No broken internal links found"])
        else:
            try:
                address_idx = gc(links_df.columns.tolist(), "Address")
            except KeyError:
                wb = Workbook()
                ws = wb.active
                if ws:
                    ws.title = "Broken Links"
                    ws.append(["Error reading broken links CSV"])
                output = io.BytesIO()
                wb.save(output)
                return output.getvalue()

            urls_with_broken_links = []
            for _, row in links_df.iterrows():
                url = str(row.iloc[address_idx])

                # Enrich
                internal_data = internal_map.get(url, {}) if internal_map else {}
                gsc_data = gsc_map.get(url, {}) if gsc_map else {}

                # Filter: indexable=True
                indexability = internal_data.get("indexability")
                if indexability != "Indexable":
                    continue

                urls_with_broken_links.append(
                    {
                        "url": url,
                        "indexability": indexability,
                        "inlinks": internal_data.get("inlinks"),
                        "impressions": gsc_data.get("impressions", 0),
                        "clicks": gsc_data.get("clicks", 0),
                    }
                )

            # Build workbook
            wb = Workbook()
            ws = wb.active
            if ws:
                ws.title = "Broken Links"
                self._write_broken_links_sheet(ws, urls_with_broken_links)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    def _write_broken_links_sheet(
        self, ws: object, urls_with_broken_links: list[dict[str, Any]]
    ) -> None:
        """Write broken links sheet."""
        ws.append([])  # type: ignore[attr-defined]
        ws.append(["NON-FUNCTIONAL INTERNAL LINKS"])  # type: ignore[attr-defined]
        ws.append([])  # type: ignore[attr-defined]

        # Summary section
        ws.append(["Summary"])  # type: ignore[attr-defined]
        ws.append(["Total Pages with Broken Links", len(urls_with_broken_links)])  # type: ignore[attr-defined]
        ws.append([])  # type: ignore[attr-defined]

        # Detailed Data section
        ws.append(["Detailed Data"])  # type: ignore[attr-defined]
        ws.append(["URL", "Inlinks", "Impressions", "Clicks"])  # type: ignore[attr-defined]

        # Sort by impressions descending
        sorted_urls = sorted(
            urls_with_broken_links, key=lambda x: x.get("impressions", 0), reverse=True
        )
        for url_data in sorted_urls:
            ws.append(  # type: ignore[attr-defined]
                [
                    safe_cell(url_data["url"]),
                    url_data.get("inlinks"),
                    url_data.get("impressions", 0),
                    url_data.get("clicks", 0),
                ]
            )
