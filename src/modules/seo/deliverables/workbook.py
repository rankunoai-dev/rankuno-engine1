"""Multi-sheet client workbook from an `AuditDataset` + `ScoringResult` (Phase 2b, ADR 0011).

Four sheets, in this order: `Overview` (per-category penalty totals),
`Issues` (one row per catalogue entry), `Pages` (the dataset's spine, themes
included), `Notes` (adapter caveats and run metadata). There is no fifth sheet
computing a single site score - `ScoringResult` does not carry one (ADR 0011
D2) and this module has no arithmetic of its own to invent one with.

**Formula injection is the load-bearing concern in this file.** openpyxl's
`Cell.value` setter classifies any string starting with `=` as a formula
(`data_type` becomes `"f"`) purely from its leading character - confirmed
against openpyxl 3.1 directly, not assumed from a style guide. A page title
or URL is content the *client's own site* produced, not Rankuno, so it is
untrusted input by the time it reaches a cell. `_safe_cell` neutralises the
leading character before any value reaches `ws.append()`, and every sheet
writer routes every text value through it - no sheet is exempted, including
`Notes`, which is closest to "operator-authored" but still not close enough
to skip the one rule that makes the rest of this file auditable.

The workbook is written with `Workbook(write_only=True)`, which restricts
every sheet to append-only rows (no cell-by-cell styling) in exchange for
bounded memory regardless of row count - the same trade `ADR 0001` already
makes for the crawl itself (20k-500k URLs, no redesign needed for the larger
path). `MAX_PAGES_PER_WORKBOOK` is the hard stop past that: a dataset over the
cap fails loud rather than shipping a client a workbook that silently dropped
rows.
"""

from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path
from typing import Final

from openpyxl import Workbook

from src.core.config import get_settings
from src.core.logger import get_logger
from src.modules.seo.contracts.audit import AuditDataset
from src.modules.seo.contracts.catalogue import ISSUE_SPECS
from src.modules.seo.deliverables.scoring import ScoringResult

__all__ = [
    "MAX_PAGES_PER_WORKBOOK",
    "SHEET_ISSUES",
    "SHEET_NOTES",
    "SHEET_OVERVIEW",
    "SHEET_PAGES",
    "WorkbookBuildError",
    "build_workbook",
]

_logger = get_logger(__name__)

MAX_PAGES_PER_WORKBOOK: Final[int] = 500_000
"""ADR 0001's stated ceiling (20k-500k URLs). A dataset past this raises
rather than truncating - a client deliverable that silently dropped rows is
exactly the RAE failure mode this package exists to close (plan §8)."""

SHEET_OVERVIEW: Final[str] = "Overview"
SHEET_ISSUES: Final[str] = "Issues"
SHEET_PAGES: Final[str] = "Pages"
SHEET_NOTES: Final[str] = "Notes"

_FORMULA_TRIGGER_CHARS: Final[frozenset[str]] = frozenset({"=", "+", "-", "@", "\t", "\r"})
"""OWASP's CSV/formula-injection trigger set. `=` is openpyxl's own trigger
for `data_type="f"`; the rest are included because a spreadsheet client may
still auto-convert them on open, and there is no cost to covering all five."""

_OVERVIEW_CAPTION: Final[str] = (
    "Per-category penalty totals. This workbook does not calculate a single "
    "site score (ADR 0011 D2)."
)


class WorkbookBuildError(ValueError):
    """`dataset` cannot become a workbook: too many pages, or the write failed.

    One exception type, matching `RulebookError` / `ScreamingFrogBundleError`:
    there is no path here on which a build failure is mistaken for success.
    """


def _safe_cell(value: object) -> object:
    """Neutralise a leading formula-trigger character before openpyxl sees it.

    Non-string values (`int`, `datetime`, `None`) pass through unchanged -
    only `str` can carry a formula-shaped payload. Prefixing with `'` costs a
    visible leading apostrophe in the rendered cell for the rare title that
    genuinely starts with one of these characters; that is the accepted
    trade-off for never writing a live formula from crawled content.
    """
    if isinstance(value, str) and value and value[0] in _FORMULA_TRIGGER_CHARS:
        return f"'{value}"
    return value


def _append_row(ws: object, values: Iterable[object]) -> None:
    """Sanitise every value in `values` and append the row to `ws`."""
    ws.append([_safe_cell(v) for v in values])  # type: ignore[attr-defined]


def _write_overview(ws: object, scoring: ScoringResult) -> None:
    _append_row(ws, (_OVERVIEW_CAPTION,))
    _append_row(ws, ())
    _append_row(
        ws,
        ("Category", "Penalty Total", "Issues Measured", "Issues Not Measured"),
    )
    for row in scoring.category_penalties:
        _append_row(
            ws,
            (
                row.category.value,
                row.penalty_total,
                row.measured_issue_count,
                row.not_measured_issue_count,
            ),
        )


def _write_issues(ws: object, scoring: ScoringResult) -> None:
    _append_row(
        ws,
        (
            "Issue ID",
            "Category",
            "Label",
            "Severity",
            "Priority",
            "Coverage",
            "Affected Pages",
            "Penalty Contribution",
        ),
    )
    for row in scoring.issue_penalties:
        spec = ISSUE_SPECS[row.issue_id]
        _append_row(
            ws,
            (
                row.issue_id.value,
                row.category.value,
                spec.label,
                spec.severity.value,
                spec.priority.value,
                row.coverage.value,
                row.affected_pages,
                row.penalty,
            ),
        )


def _write_pages(ws: object, dataset: AuditDataset) -> None:
    _append_row(ws, ("URL", "Theme 1", "Theme 2", "Language", "Business Priority"))
    for page in sorted(dataset.pages, key=lambda p: p.url):
        _append_row(
            ws,
            (page.url, page.theme_1, page.theme_2, page.language, page.business_priority),
        )


def _write_notes(ws: object, dataset: AuditDataset) -> None:
    _append_row(ws, ("Site", dataset.site))
    _append_row(ws, ("Produced At (UTC)", dataset.produced_at.isoformat()))
    _append_row(ws, ("Source", dataset.source.value))
    _append_row(ws, ())
    _append_row(ws, ("Notes",))
    for note in dataset.notes:
        _append_row(ws, (note,))


def build_workbook(
    dataset: AuditDataset, scoring: ScoringResult, *, output_dir: Path | None = None
) -> Path:
    """Write the four-sheet client workbook for `dataset` and return its path.

    Args:
        dataset: The dataset to render. `dataset.pages` should already carry
            themes from `apply_rulebook` if a rulebook was applied; this
            function does not classify anything itself.
        scoring: The `score_dataset(dataset)` result. Not recomputed here so a
            caller's chosen `weights=` is respected without a second pass.
        output_dir: Directory to write into. Defaults to
            `get_settings().deliverables_output_dir`.

    Returns:
        Path to the written `.xlsx` file:
        `{output_dir}/{dataset.site}-{produced_at:%Y%m%dT%H%M%SZ}.xlsx`.

    Raises:
        WorkbookBuildError: `dataset.pages` exceeds `MAX_PAGES_PER_WORKBOOK`,
            or the write itself failed.
    """
    if len(dataset.pages) > MAX_PAGES_PER_WORKBOOK:
        msg = (
            f"{dataset.site}: {len(dataset.pages)} pages exceeds "
            f"MAX_PAGES_PER_WORKBOOK ({MAX_PAGES_PER_WORKBOOK}); refusing to "
            "build a workbook that would have to truncate rows"
        )
        raise WorkbookBuildError(msg)

    target_dir = output_dir if output_dir is not None else get_settings().deliverables_output_dir
    target_dir.mkdir(parents=True, exist_ok=True)
    path = target_dir / f"{dataset.site}-{dataset.produced_at:%Y%m%dT%H%M%SZ}.xlsx"

    workbook = Workbook(write_only=True)
    _write_overview(workbook.create_sheet(SHEET_OVERVIEW), scoring)
    _write_issues(workbook.create_sheet(SHEET_ISSUES), scoring)
    _write_pages(workbook.create_sheet(SHEET_PAGES), dataset)
    _write_notes(workbook.create_sheet(SHEET_NOTES), dataset)

    try:
        workbook.save(path)
    except OSError as exc:
        raise WorkbookBuildError(f"{path}: failed to write workbook: {exc}") from exc

    _logger.info(
        "workbook_built",
        extra={"site": dataset.site, "pages": len(dataset.pages), "path": str(path)},
    )
    return path
