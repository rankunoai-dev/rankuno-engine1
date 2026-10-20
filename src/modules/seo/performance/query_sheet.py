"""How the top queries and the partial-data banner appear in the URL exports.

What this module decides is presentation: the cell text, the banner wording, and
the layout of the long-format sheet. Which queries are "top" and which URL each
belongs to is decided upstream, in `query_by_url.py`, so this module never has
to know how a query was matched.

It lives in `performance` because it renders Search Console data, and the
`page_classifier` URL export imports it. That import is allowed; the seam ADR 0011
keeps closed is between `page_classifier` and `deliverables`.

Every text value written here goes through `safe_cell`. A query is typed by a
searcher and a URL is set by the site, so both are untrusted by the time they
reach a cell, and a spreadsheet will run a leading `=` as a formula.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from typing import Final

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from src.integrations.gsc_schemas import GscPageQueries, GscQueryRow
from src.modules.seo.contracts.cell_guard import safe_cell

__all__ = [
    "CATEGORY_HEADER",
    "GSC_COMPLETE_NOTE",
    "GSC_PARTIAL_PREFIX",
    "TOP_QUERIES_HEADER",
    "TOP_QUERIES_SHEET",
    "partial_banner",
    "top_queries_cell",
    "write_top_queries_sheet",
]

TOP_QUERIES_HEADER: Final[str] = "Top Queries"
"""The "All URLs" column that holds each page's top queries in one cell."""

CATEGORY_HEADER: Final[str] = "Category"
"""The "All URLs" column holding the page's path category (`contracts.url_category`)."""

TOP_QUERIES_SHEET: Final[str] = "Top Queries"
"""The long-format sheet: one row per URL and rank, for filtering and pivoting."""

GSC_PARTIAL_PREFIX: Final[str] = "GSC data partial"
"""The words an analyst searches for. Every partial banner begins with them."""

GSC_COMPLETE_NOTE: Final[str] = "GSC data complete"

_SHEET_HEADERS: Final[tuple[str, ...]] = (
    "URL",
    "Rank",
    "Query",
    "Clicks",
    "Impressions",
    "Position",
)
_HEADER_FILL: Final = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
_HEADER_FONT: Final = Font(bold=True, color="FFFFFF")


def top_queries_cell(rows: Sequence[GscQueryRow]) -> str | None:
    """A URL's top queries as one cell: `query (clicks)` joined by `; `.

    Args:
        rows: The URL's ranked queries, already limited to the top five.

    Returns:
        The cell text, or `None` (a blank cell) when the URL has no queries. A
        blank is never a zero, because an unfetched URL and a URL with no queries
        look the same in a spreadsheet.
    """
    if not rows:
        return None
    text = "; ".join(f"{row.query} ({row.clicks})" for row in rows)
    # `safe_cell` is typed for any cell value; for a string it returns a string.
    return str(safe_cell(text))


def partial_banner(*, page_metrics_complete: bool, queries: GscPageQueries | None) -> str | None:
    """The "GSC data partial" line, or `None` when every fetch was complete.

    A partial fetch keeps the rows that arrived, so the export still carries them.
    The banner exists so a reader does not take a short total for a true one.

    Args:
        page_metrics_complete: Whether the page-level fetch reached its last page.
        queries: The query-by-page fetch, or `None` if none was taken.

    Returns:
        The banner text, or `None`.
    """
    parts: list[str] = []
    if not page_metrics_complete:
        parts.append("page metrics")
    if queries is not None and not queries.complete:
        parts.append(f"top queries ({queries.stop_reason})")
    if not parts:
        return None
    return (
        f"{GSC_PARTIAL_PREFIX}: Search Console did not return every row for "
        f"{' and '.join(parts)}. Rows that did arrive are included, so GSC "
        "figures may be understated or incomplete."
    )


def write_top_queries_sheet(
    ws: Worksheet,
    urls: Sequence[str],
    rows_by_url: Mapping[str, Sequence[GscQueryRow]],
    banner: str | None,
) -> None:
    """Write the long-format Top Queries sheet.

    Row 1 is the status line, which is the banner when data is partial and
    "GSC data complete" otherwise. It is always present, so an empty sheet cannot
    be mistaken for a complete one with nothing to report. Row 2 is the header.

    Args:
        ws: The sheet to write into.
        urls: Crawled URLs in the order the sheet should list them.
        rows_by_url: Each URL's ranked queries.
        banner: The partial banner from `partial_banner`, or `None`.
    """
    ws.append([safe_cell(banner or GSC_COMPLETE_NOTE)])
    ws.append(list(_SHEET_HEADERS))
    for cell in ws[2]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for url in urls:
        for rank, row in enumerate(rows_by_url.get(url, ()), start=1):
            ws.append(
                [
                    safe_cell(url),
                    rank,
                    safe_cell(row.query),
                    row.clicks,
                    row.impressions,
                    row.avg_position,
                ]
            )
